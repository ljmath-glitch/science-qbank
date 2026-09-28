"""Rebuild the supplied A4 mock-exam cover with nine editable Word boxes.

All fixed artwork is retained at the source PDF's physical coordinates as a
300-dpi page background.  The exam label, scope and seven grade thresholds are
native Word VML text boxes.  This intentionally does not imply that the other
printed instructions are editable.
"""

from __future__ import annotations

import argparse
import copy
import subprocess
from pathlib import Path
from xml.sax.saxutils import escape

import pdfplumber
from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
V_NS = "urn:schemas-microsoft-com:vml"
O_NS = "urn:schemas-microsoft-com:office:office"

DEFAULT_RANGES = ["48~50", "46~47", "43~45", "36~42", "29~35", "18~28", "0~17"]


def emu(pt: float) -> int:
    return round(pt * 12700)


def make_anchor(inline):
    anchor = OxmlElement("wp:anchor")
    for attr, value in {
        "distT": "0", "distB": "0", "distL": "0", "distR": "0",
        "simplePos": "0", "relativeHeight": "0", "behindDoc": "1",
        "locked": "1", "layoutInCell": "1", "allowOverlap": "1",
    }.items():
        anchor.set(attr, value)
    anchor.append(parse_xml('<wp:simplePos xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" x="0" y="0"/>'))
    for axis in ("H", "V"):
        position = OxmlElement(f"wp:position{axis}")
        position.set("relativeFrom", "page")
        offset = OxmlElement("wp:posOffset")
        offset.text = "0"
        position.append(offset)
        anchor.append(position)
    for tag in ("extent", "effectExtent"):
        element = inline.find(qn(f"wp:{tag}"))
        if element is not None:
            anchor.append(copy.deepcopy(element))
    anchor.append(OxmlElement("wp:wrapNone"))
    for tag in ("docPr", "cNvGraphicFramePr"):
        element = inline.find(qn(f"wp:{tag}"))
        if element is not None:
            anchor.append(copy.deepcopy(element))
    anchor.append(copy.deepcopy(inline.find(qn("a:graphic"))))
    return anchor


def textbox(name, text, x, y, width, height, fill, font, font_size, top_inset, bold=False):
    # The PDF coordinate origin is the same top-left page origin used here.
    shape_style = (
        f"position:absolute;left:{x:.3f}pt;top:{y:.3f}pt;"
        f"width:{width:.3f}pt;height:{height:.3f}pt;z-index:251658240;"
        "mso-position-horizontal:absolute;mso-position-horizontal-relative:page;"
        "mso-position-vertical:absolute;mso-position-vertical-relative:page;"
        "v-text-anchor:middle"
    )
    weight = '<w:b/>' if bold else ""
    xml = f'''<w:pict xmlns:w="{W_NS}" xmlns:v="{V_NS}" xmlns:o="{O_NS}">
      <v:rect id="{escape(name)}" style="{shape_style}" fillcolor="{fill}" stroked="f" o:allowincell="f">
        <v:textbox inset="0,{top_inset:.2f}pt,0,0" style="mso-fit-shape-to-text:false">
          <w:txbxContent><w:p><w:pPr><w:jc w:val="center"/><w:spacing w:before="0" w:after="0" w:line="{round(font_size * 20)}" w:lineRule="exact"/></w:pPr>
            <w:r><w:rPr><w:rFonts w:ascii="{escape(font)}" w:hAnsi="{escape(font)}" w:eastAsia="{escape(font)}"/>
              <w:sz w:val="{round(font_size * 2)}"/>{weight}</w:rPr><w:t>{escape(text)}</w:t></w:r>
          </w:p></w:txbxContent>
        </v:textbox>
      </v:rect>
    </w:pict>'''
    return parse_xml(xml)


def build(args):
    source = Path(args.source)
    output = Path(args.output)
    if output.exists():
        raise FileExistsError(f"Refusing to overwrite: {output}")
    ranges = args.ranges.split(",") if args.ranges else DEFAULT_RANGES
    if len(ranges) != 7 or any(not value.strip() for value in ranges):
        raise ValueError("--ranges must contain seven non-empty comma-separated values")
    with pdfplumber.open(source) as pdf:
        if len(pdf.pages) != 1:
            raise ValueError("The cover source must contain exactly one page")
        page = pdf.pages[0]
        page_width, page_height = float(page.width), float(page.height)
        cells = sorted(
            (r for r in page.rects
             if 706.5 < float(r["top"]) < 708 and 19 < float(r["bottom"] - r["top"]) < 20
             and 57 < float(r["x1"] - r["x0"]) < 59),
            key=lambda r: float(r["x0"]),
        )
        if len(cells) != 7:
            raise ValueError(f"Expected seven measured grade cells, found {len(cells)}")
    background = Path(args.workdir) / "mock-cover-background.png"
    subprocess.run([
        args.pdftoppm, "-f", "1", "-singlefile", "-r", "300", "-png",
        str(source), str(background.with_suffix("")),
    ], check=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Pt(page_width), Pt(page_height)
    section.top_margin = section.bottom_margin = Pt(0)
    section.left_margin = section.right_margin = Pt(0)
    section.header_distance = section.footer_distance = Pt(0)
    normal = doc.styles["Normal"]
    normal.paragraph_format.space_before = normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = Pt(1)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    image = paragraph.add_run().add_picture(str(background), width=Pt(page_width), height=Pt(page_height))
    inline = image._inline
    inline.getparent().replace(inline, make_anchor(inline))
    # Current PDF has visible FreeText annotations over old fixed text.  Opaque
    # white Word boxes cover both; the warning frame below starts at 167.4 pt.
    fields = [
        ("exam_round", args.exam, 251.0, 144.0, 107.5, 20.5, "#FFFFFF", "PingFang TC", 12.5, 8.5, True),
        ("book_range", args.scope, 453.0, 139.0, 87.0, 25.0, "#FFFFFF", "PingFang TC", 16.0, 4.0, False),
    ]
    for index, (cell, value) in enumerate(zip(cells, ranges), 1):
        fields.append((
            f"correct_answers_{index}", value,
            float(cell["x0"]), float(cell["top"]),
            float(cell["x1"] - cell["x0"]), float(cell["bottom"] - cell["top"]),
            "#E5E4E4", "Arial Black", 9.0, 4.2, False,
        ))
    run = paragraph.add_run()
    for field in fields:
        run._r.append(textbox(*field))
    doc.core_properties.title = "茲茲文教自然模考封面"
    doc.core_properties.subject = "考次、範圍與七段答對題數可編輯"
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--workdir", required=True)
    parser.add_argument("--pdftoppm", required=True)
    parser.add_argument("--exam", default="【114國九二模】")
    parser.add_argument("--scope", default="第１~４冊")
    parser.add_argument("--ranges", default="")
    build(parser.parse_args())
